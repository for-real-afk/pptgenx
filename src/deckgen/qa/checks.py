"""Cheap automated checks on generated shapes (python-pptx). Catches what is geometric and certain;
the visual critic covers what needs eyes."""
from __future__ import annotations

from dataclasses import dataclass
from itertools import combinations
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deckgen.config import CONFIG
from deckgen.schemas.qa import Issue

EMU = 914400
SLIDE_W, CONTENT_BOTTOM = 13.333, 6.5
AVG_CHAR_EM, LINE_EM = 0.5, 1.2
BENIGN_OVERLAP = ("Timeline rail", " marker")   # rail passes through its markers by design
CAPTION_OK = ("Source caption", "Chart caption")
BALANCED_LAYOUTS = {"TITLE_ONLY", "CHART_COMMENTARY", "TABLE", "PHOTO_TEXT"}


@dataclass(frozen=True)
class Box:
    x: float
    y: float
    w: float
    h: float

    @property
    def r(self) -> float: return self.x + self.w
    @property
    def b(self) -> float: return self.y + self.h

    def inter(self, o: Box) -> float:
        iw, ih = min(self.r, o.r) - max(self.x, o.x), min(self.b, o.b) - max(self.y, o.y)
        return iw * ih if iw > 0 and ih > 0 else 0.0

    def contains(self, o: Box, tol: float = 0.03) -> bool:
        return o.x >= self.x - tol and o.y >= self.y - tol and o.r <= self.r + tol and o.b <= self.b + tol


def _box(sh) -> Box:
    return Box(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)


def estimate_text_height(sh) -> float | None:
    """Inches of height the text needs at its explicit font sizes; None if sizes are unknown."""
    tf = sh.text_frame
    usable = sh.width / EMU - (tf.margin_left + tf.margin_right) / EMU
    total = 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        if not sizes:
            return None
        size = max(sizes)
        indent = 0.3 if p._p.xpath("./a:pPr/a:buChar|./a:pPr/a:buAutoNum") else 0.0
        width = max(usable - indent, 0.3)
        chars_per_line = max(1, int(width * 72 / (size * AVG_CHAR_EM)))
        words, lines, cur = text.split(), 1, 0
        for w in words:  # greedy word wrap
            n = len(w) + (1 if cur else 0)
            if cur + n > chars_per_line and cur:
                lines, cur = lines + 1, len(w)
            else:
                cur += n
        after = p.space_after.pt / 72 if p.space_after else 0.0
        total += lines * size * LINE_EM / 72 + after
    return total + (tf.margin_top + tf.margin_bottom) / EMU


def check(pptx: Path) -> list[Issue]:
    issues: list[Issue] = []
    for n, slide in enumerate(Presentation(pptx).slides, 1):
        def add(t, sev, desc, fix, n=n):
            issues.append(Issue(slide=n, type=t, severity=sev, description=desc, suggested_fix=fix, source="auto"))

        shapes = [s for s in slide.shapes
                  if not (s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)]
        boxes = {s.shape_id: _box(s) for s in shapes}
        for s in shapes:
            b, name = boxes[s.shape_id], s.name
            if b.x < -0.01 or b.y < -0.01 or b.r > SLIDE_W + 0.01 or b.b > 7.51:
                add("off_slide", "high", f"'{name}' extends past the slide edge", "resize or reposition inside the slide")
            elif b.x < CONFIG.min_margin_in - 0.01 or b.r > SLIDE_W - CONFIG.min_margin_in + 0.01:
                add("margin", "medium", f"'{name}' is inside the {CONFIG.min_margin_in} in margin", "move inward")
            if s.has_text_frame:
                text = s.text_frame.text.strip()
                if not text and s.is_placeholder:
                    add("placeholder_text", "high", f"empty placeholder '{name}'", "fill or remove it")
                need = estimate_text_height(s) if text else None
                if need and need > b.h * 1.02:
                    add("overflow", "high", f"text in '{name}' needs ~{need:.2f} in but box is {b.h:.2f} in", "shorten the text or enlarge the container")
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if not r.font.size:
                            continue
                        pt = r.font.size.pt
                        if pt < CONFIG.min_caption_pt:
                            add("font_floor", "high", f"{pt:g} pt text in '{name}' is below the {CONFIG.min_caption_pt} pt floor", "cut words or enlarge the container")
                        elif pt < CONFIG.min_body_pt and name not in CAPTION_OK and slide.slide_layout.name != "SOURCES":
                            add("font_floor", "medium", f"{pt:g} pt body text in '{name}' is below {CONFIG.min_body_pt} pt", "cut words or enlarge the container")
        for a, c in combinations(shapes, 2):
            ba, bc = boxes[a.shape_id], boxes[c.shape_id]
            if any(k in a.name or k in c.name for k in BENIGN_OVERLAP):
                continue
            area = ba.inter(bc)
            if area > 0.02 and not ba.contains(bc) and not bc.contains(ba):
                add("overlap", "high", f"'{a.name}' overlaps '{c.name}' ({area:.2f} sq in)", "separate the elements")
        if slide.slide_layout.name in BALANCED_LAYOUTS:
            body = [boxes[s.shape_id] for s in shapes if s.name not in CAPTION_OK and 1.6 <= boxes[s.shape_id].y < 6.6]
            if body and CONTENT_BOTTOM - max(b.b for b in body) > 1.0:
                gap = CONTENT_BOTTOM - max(b.b for b in body)
                add("empty_space", "medium", f"content stops {gap:.1f} in above the bottom of the content area", "grow the content to fill the area or rebalance vertically")
    return issues
