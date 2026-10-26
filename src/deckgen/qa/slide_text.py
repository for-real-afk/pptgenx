"""Content QA: extract every slide's title, text, tables, chart data, alt text and notes."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deckgen.schemas.qa import Issue

LEFTOVER = re.compile(r"\bxxx+\b|lorem|ipsum|\bTODO\b|\[insert|this page layout|click to edit", re.I)
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


@dataclass(frozen=True)
class SlideText:
    number: int
    layout: str
    title: str
    texts: tuple[str, ...]
    tables: tuple[tuple[tuple[str, ...], ...], ...]
    charts: tuple[str, ...]
    alt_texts: tuple[str, ...]
    notes: str


def _shape_title(shape) -> bool:
    # pptxgenjs placeholders have idx != 0, so python-pptx's shapes.title is None: match by type.
    return shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES


def extract(pptx: Path) -> list[SlideText]:
    out = []
    for i, slide in enumerate(Presentation(pptx).slides, 1):
        title, texts, tables, charts, alts = "", [], [], [], []
        for sh in slide.shapes:
            descr = sh._element.xpath(".//p:cNvPr/@descr")
            if descr and descr[0]:
                alts.append(descr[0])
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = sh.text_frame.text.strip()
                if _shape_title(sh):
                    title = txt
                else:
                    texts.append(txt)
            if getattr(sh, "has_table", False) and sh.has_table:
                tables.append(tuple(tuple(c.text for c in r.cells) for r in sh.table.rows))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                cats = list(ch.plots[0].categories)
                charts.append("; ".join(f"{s.name}: " + ", ".join(f"{c}={v}" for c, v in zip(cats, s.values)) for s in ch.plots[0].series))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append(SlideText(i, slide.slide_layout.name, title, tuple(texts), tuple(tables), tuple(charts), tuple(alts), notes))
    return out


def to_markdown(slides: list[SlideText]) -> str:
    md = []
    for s in slides:
        md.append(f"## Slide {s.number} [{s.layout}] {s.title}")
        md += [f"- {t}" for t in s.texts]
        md += [f"- table: {t}" for t in s.tables] + [f"- chart: {c}" for c in s.charts]
        md += [f"- alt: {a}" for a in s.alt_texts]
        md.append(f"- notes: {s.notes or '(none)'}\n")
    return "\n".join(md)


def content_issues(slides: list[SlideText]) -> list[Issue]:
    issues = []
    for s in slides:
        blob = " ".join([s.title, *s.texts, s.notes])
        if not s.title:
            issues.append(Issue(slide=s.number, type="other", severity="high", description="slide has no title", suggested_fix="add a message-stating title", source="content"))
        if m := LEFTOVER.search(blob):
            issues.append(Issue(slide=s.number, type="placeholder_text", severity="high", description=f"leftover placeholder text '{m.group(0)}'", suggested_fix="replace with real content", source="content"))
        if s.title.endswith("."):
            issues.append(Issue(slide=s.number, type="other", severity="low", description="title ends with a period", suggested_fix="drop the period", source="content"))
    return issues
