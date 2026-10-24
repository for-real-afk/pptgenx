"""QA tests. The broken-slide deck must be caught by BOTH the automated checks and the live critic."""
import json
import os
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from deckgen.llm import FakeLLM
from deckgen.qa import checks, soffice, validate_pptx
from deckgen.qa.slide_text import content_issues, extract
from deckgen.qa.visual_critic import critique
from deckgen.schemas.qa import CriticReport, Issue

ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "brand" / "default_v1"
FIXTURE = ROOT / "tests" / "fixtures" / "handwritten_spec.json"
LIVE = pytest.mark.skipif(not os.environ.get("DECKGEN_LIVE") or not os.environ.get("GEMINI_API_KEY"),
                          reason="set DECKGEN_LIVE=1 and GEMINI_API_KEY for live critic tests")


@pytest.fixture(scope="module")
def good_deck(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("good") / "good.pptx"
    r = subprocess.run(["node", "renderer/build_deck.js", str(FIXTURE), str(BRAND), str(out)], cwd=ROOT, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return out


@pytest.fixture(scope="module")
def broken_deck(good_deck, tmp_path_factory) -> Path:
    """Slide 3 gets a tiny box with a wall of 14 pt text (overflow) and a rectangle straddling a card (overlap)."""
    prs = Presentation(good_deck)
    s = prs.slides[2]
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(2.0), Inches(0.5))
    tb.name = "Broken overflow box"
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "This sentence is far too long for the little box it was put in. " * 4
    run.font.size = Pt(14)
    bar = s.shapes.add_shape(1, Inches(4.2), Inches(2.5), Inches(3.0), Inches(1.5))
    bar.name = "Broken overlap rect"
    out = tmp_path_factory.mktemp("bad") / "broken.pptx"
    prs.save(out)
    return out


def test_good_deck_has_no_auto_or_file_or_content_issues(good_deck):
    assert checks.check(good_deck) == []
    assert validate_pptx.validate(good_deck) == []
    assert content_issues(extract(good_deck)) == []


def test_auto_checks_catch_overflow_and_overlap(broken_deck):
    found = {(i.slide, i.type) for i in checks.check(broken_deck)}
    assert (3, "overflow") in found
    assert (3, "overlap") in found


def test_file_validator_catches_corrupt_deck(good_deck, tmp_path):
    bad = tmp_path / "corrupt.pptx"
    with zipfile.ZipFile(good_deck) as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/charts/chart"):
                data = data.replace(b'<a:srgbClr val="', b'<a:srgbClr val="tx1', 1)  # scheme name in a hex slot
            if item.filename == "ppt/slides/slide3.xml":
                data = data.replace(b"</p:sld>", b"<broken></p:sld>")                # not well-formed
            zout.writestr(item, data)
    errs = validate_pptx.validate(bad)
    assert any("not well-formed" in e for e in errs)
    assert any("invalid color" in e for e in errs)


def test_leftover_placeholder_text_flagged(good_deck, tmp_path):
    prs = Presentation(good_deck)
    prs.slides[0].shapes[1].text_frame.text = "Lorem ipsum dolor"
    p = tmp_path / "lorem.pptx"
    prs.save(p)
    assert any(i.type == "placeholder_text" for i in content_issues(extract(p)))


def test_critic_plumbing_with_fake_llm(tmp_path):
    from PIL import Image

    imgs = []
    for k in range(7):
        p = tmp_path / f"slide-{k + 1}.jpg"
        Image.new("RGB", (1600, 900), "white").save(p)
        imgs.append(p)

    def script(prompt, images):
        first = int(prompt.split("slides [")[1].split("]")[0].split(",")[0])
        return CriticReport(issues=[Issue(slide=first, type="overflow", severity="high", description="x", suggested_fix="y"),
                                    Issue(slide=99, type="overlap", severity="low", description="out of range", suggested_fix="z")])

    llm = FakeLLM({"critic": script, "critic_consistency": CriticReport(issues=[])})
    issues = critique(llm, imgs, list(range(1, 8)), all_images=imgs)
    assert sorted(i.slide for i in issues) == [1, 6]  # one per batch of 5; slide 99 dropped
    assert [c["stage"] for c in llm.calls] == ["critic", "critic", "critic_consistency"]
    assert all(c["images"] for c in llm.calls)
    assert not any(w in c["prompt"] + c["system"] for c in llm.calls for w in ("DeckSpec", "SlideSpec", "pptxgenjs", "layouts.js", "placeholders"))  # critic is isolated


@LIVE
def test_live_critic_flags_broken_slide(broken_deck, tmp_path):
    try:
        soffice.backend()
    except soffice.RenderError:
        pytest.skip("no render backend")
    from deckgen.llm import GeminiLLM

    imgs = soffice.render_deck(broken_deck, tmp_path)
    issues = critique(GeminiLLM(), [imgs[2]], [3])
    kinds = {i.type for i in issues if i.slide == 3 and i.severity in ("high", "medium")}
    assert kinds & {"overflow", "overlap"}, issues


@LIVE
def test_recorded_empty_space_before_after():
    """runs/critic_before.json / critic_after.json: live critic runs around the compose fix."""
    before = json.loads((ROOT / "runs" / "critic_before.json").read_text())
    after = json.loads((ROOT / "runs" / "critic_after.json").read_text())
    assert any(i["slide"] == 6 and i["type"] == "empty_space" for i in before)
    assert not any(i["slide"] == 6 for i in after)
