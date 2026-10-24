import json
import subprocess
from pathlib import Path

import pytest

from deckgen.schemas import DeckSpec
from deckgen.spec.validate_spec import load_manifest, validate_spec

ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "brand" / "default_v1"
FIXTURE = ROOT / "tests" / "fixtures" / "handwritten_spec.json"


def node(*args):
    return subprocess.run(["node", *args], capture_output=True, text=True, cwd=ROOT)


def load_fixture() -> DeckSpec:
    return DeckSpec.model_validate_json(FIXTURE.read_text(encoding="utf-8"))


def test_grid_math():
    js = ("const g=require('./renderer/grid');const c=g.cols(3);"
          "console.log(JSON.stringify([c.length, +(c[0].w*3+g.GAP*2).toFixed(3), +g.CONTENT.w.toFixed(3), g.grid(2,3).length]))")
    assert json.loads(node("-e", js).stdout) == [3, 12.133, 12.133, 6]


def test_fixture_spec_valid():
    node("renderer/gen_manifest.js", str(BRAND))
    assert validate_spec(load_fixture(), load_manifest(BRAND)) == []


def test_bad_layout_placeholder_and_capacity_caught():
    spec = load_fixture()
    spec.slides[0].layout = "NOPE"
    spec.slides[1].placeholders = {"bogus": "x"}
    spec.slides[8].placeholders = {"body": ["word " * 400]}
    errs = validate_spec(spec, load_manifest(BRAND))
    assert any("NOPE" in e for e in errs)
    assert any("bogus" in e for e in errs)
    assert any("allows" in e and "s09" in e for e in errs)


def test_renderer_rejects_misspelt_layout(tmp_path):
    bad = json.loads(FIXTURE.read_text(encoding="utf-8"))
    bad["slides"][0]["layout"] = "TITEL"
    p = tmp_path / "bad.json"
    p.write_text(json.dumps(bad))
    r = node("renderer/build_deck.js", str(p), str(BRAND), str(tmp_path / "x.pptx"))
    assert r.returncode != 0 and "TITEL" in r.stderr


def test_apply_theme_throws_on_scheme_in_hex_option(tmp_path):
    js = """
    const P=require('pptxgenjs');const {applyTheme}=require('./renderer/apply_theme');
    (async()=>{const p=new P();const s=p.addSlide();
    s.addChart(p.charts.BAR,[{name:'a',labels:['x','y'],values:[1,2]}],{x:1,y:1,w:5,h:3,valGridLine:{color:'tx1',size:1}});
    await p.writeFile({fileName:process.argv[1]});
    const theme=JSON.parse(require('fs').readFileSync('brand/default_v1/theme.json'));
    try{await applyTheme(process.argv[1],theme);console.log('NO_THROW')}catch(e){console.log('THREW')}})()
    """
    assert "THREW" in node("-e", js, str(tmp_path / "c.pptx")).stdout


def test_build_fixture_deck_opens(tmp_path):
    from pptx import Presentation

    out = tmp_path / "f.pptx"
    r = node("renderer/build_deck.js", str(FIXTURE), str(BRAND), str(out))
    assert r.returncode == 0, r.stderr
    assert len(Presentation(out).slides) == 10


def test_fixture_render_one_image_per_slide(tmp_path):
    from deckgen.qa import soffice

    try:
        soffice.backend()
    except soffice.RenderError:
        pytest.skip("no render backend")
    out = tmp_path / "f.pptx"
    assert node("renderer/build_deck.js", str(FIXTURE), str(BRAND), str(out)).returncode == 0
    assert len(soffice.render_deck(out, tmp_path / "r")) == 10
